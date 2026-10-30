#!/usr/bin/env python3
"""Build a clearly labelled partial Arabic guide from the checked-in Markdown.

Requires python-docx, reportlab, arabic-reshaper, and python-bidi. The output
files are created only when absent; this script never overwrites a deliverable.
The screenshots are already in screenshots/raw and are not generated here.
"""
from __future__ import annotations

import html
import io
import re
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, KeepTogether, PageBreak, PageTemplate, Spacer, Flowable

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'USER-GUIDE.ar.md'
OUT = ROOT / 'deliverables'
PDF = OUT / 'KAYAN-User-Guide-AR-DRAFT.pdf'
DOCX = OUT / 'KAYAN-User-Guide-AR-DRAFT.docx'
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
FONT_BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
BASE_SHA = '320eab32c8debfe39a3fc75bbf9ebaebb189f5b1'


def clean(text: str) -> str:
    text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
    # Isolate embedded Latin paths/codes from the surrounding RTL paragraph.
    text = re.sub(r'`([^`]+)`', lambda m: '\u200e' + m.group(1) + '\u200e', text)
    return text.replace('**', '').strip()


def visual(text: str) -> str:
    return get_display(arabic_reshaper.reshape(text))


def wrap(text: str, font: str, size: float, width: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ''
    for word in words:
        while pdfmetrics.stringWidth(visual(word), font, size) > width and len(word) > 1:
            if line:
                lines.append(line)
                line = ''
            cut = len(word) - 1
            while cut > 1 and pdfmetrics.stringWidth(visual(word[:cut]), font, size) > width:
                cut -= 1
            lines.append(word[:cut])
            word = word[cut:]
        candidate = (line + ' ' + word).strip()
        if line and pdfmetrics.stringWidth(visual(candidate), font, size) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    return lines or ['']


class RLine(Flowable):
    def __init__(self, text: str, *, size=10.7, leading=17, bold=False, color=colors.HexColor('#203044'), chapter=None):
        super().__init__()
        self.text, self.size, self.leading = text, size, leading
        self.font = 'KayanBold' if bold else 'Kayan'
        self.color, self.chapter = color, chapter
        self.height = leading
        if chapter is not None:
            self.keepWithNext = True

    def wrap(self, availWidth, availHeight):
        self.width = availWidth
        return availWidth, self.leading

    def draw(self):
        self.canv.setFont(self.font, self.size)
        self.canv.setFillColor(self.color)
        self.canv.drawRightString(self.width, self.leading - self.size - 1, visual(self.text))


class GuideDoc(BaseDocTemplate):
    def __init__(self, target):
        super().__init__(target, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                         topMargin=20*mm, bottomMargin=19*mm, title='KAYAN ERP Arabic user guide — draft')
        self.chapters = {}
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates(PageTemplate(id='guide', frames=[frame], onPage=self.decorate))

    def decorate(self, canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#DCE5EF'))
        canvas.line(self.leftMargin, 15*mm, A4[0] - self.rightMargin, 15*mm)
        canvas.setFont('Kayan', 8.5)
        canvas.setFillColor(colors.HexColor('#66788A'))
        canvas.drawRightString(A4[0] - self.rightMargin, 10*mm, visual('دليل كيان — مسودة جزئية'))
        canvas.drawString(self.leftMargin, 10*mm, str(doc.page))
        canvas.restoreState()

    def afterFlowable(self, f):
        if getattr(f, 'chapter', None) is not None:
            self.chapters[f.chapter] = self.page


def parse_source():
    text = SOURCE.read_text(encoding='utf-8')
    sections = re.findall(r'^## (\d+)\. (.+)$', text, re.M)
    if [int(n) for n, _ in sections] != list(range(1, 21)):
        raise ValueError('Expected exactly 20 numbered sections in order')
    content = text[text.index('## 1. '):]
    return [(int(n), title) for n, title in sections], content.splitlines()


def img_path(raw: str) -> Path:
    path = (ROOT / raw).resolve()
    if not path.is_relative_to((ROOT / 'screenshots' / 'raw').resolve()) or not path.is_file():
        raise ValueError(f'Invalid or missing screenshot: {raw}')
    return path


def block(text, *, width, size=10.7, leading=17, bold=False, color=colors.HexColor('#203044'), chapter=None):
    font = 'KayanBold' if bold else 'Kayan'
    return [RLine(line, size=size, leading=leading, bold=bold, color=color, chapter=chapter if i == 0 else None)
            for i, line in enumerate(wrap(clean(text), font, size, width))]


def pdf_story(sections, lines, pages):
    w = A4[0] - 40*mm
    story = [Spacer(1, 55*mm)]
    story += block('دليل استخدام نظام كيان ERP', width=w, size=24, leading=37, bold=True, color=colors.HexColor('#14558B'))
    story += block('مسودة عربية مصوّرة جزئيًا — للمراجعة', width=w, size=15, leading=25, bold=True)
    story += [Spacer(1, 13*mm)]
    story += block('الصور المحفوظة من واجهة سابقة؛ باقي اللقطات مطلوبة. لم تختبر خطوات البرنامج عمليًا في هذا العمل.', width=w, size=11, leading=20)
    story += block('المرجع: شجرة العمل على الفرع arena/2d794f08-erp-kayan، مع تغييرات محلية غير مدمجة.', width=w, size=10, leading=18)
    story += block('HEAD: ' + BASE_SHA, width=w, size=9, leading=16)
    story += [PageBreak()]
    story += block('فهرس الأقسام', width=w, size=18, leading=29, bold=True, color=colors.HexColor('#14558B'))
    story += [Spacer(1, 4*mm)]
    for n, title in sections:
        page = pages.get(n, 0)
        # Explicit page value only after a pagination pass.
        story += block(f'{n}. {title}   —   {page}', width=w, size=10.2, leading=19)
    idx = 0
    while idx < len(lines):
        raw = lines[idx].strip()
        idx += 1
        if not raw or raw == '---':
            if raw == '---': story.append(Spacer(1, 4*mm))
            continue
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', raw)
        if image:
            path = img_path(image.group(2))
            im = Image(str(path))
            original_width, original_height = im.imageWidth, im.imageHeight
            im.drawWidth = w
            im.drawHeight = w * original_height / original_width
            caption = []
            while idx < len(lines) and not lines[idx].strip(): idx += 1
            if idx < len(lines) and lines[idx].startswith('**الصورة'):
                caption = block(lines[idx], width=w, size=9.3, leading=15, bold=True)
                idx += 1
            story.append(KeepTogether([im, Spacer(1, 3*mm), *caption, Spacer(1, 4*mm)]))
            continue
        head = re.match(r'^## (\d+)\. (.+)', raw)
        if head:
            story += [PageBreak()]
            story += block(f'{head.group(1)}. {head.group(2)}', width=w, size=16, leading=26, bold=True,
                           color=colors.HexColor('#14558B'), chapter=int(head.group(1)))
            story += [Spacer(1, 2*mm)]
            continue
        if raw.startswith('### '):
            story += [Spacer(1, 2*mm)]
            story += block(raw[4:], width=w, size=12.3, leading=21, bold=True, color=colors.HexColor('#176F77'))
            continue
        if re.match(r'^\|[\s:|-]+\|$', raw):
            continue
        if raw.startswith('|'):
            raw = ' — '.join(x.strip() for x in raw.strip('|').split('|') if x.strip())
        elif raw.startswith('>'):
            raw = raw[1:].strip()
        story += block(raw, width=w)
        story += [Spacer(1, 2*mm)]
    return story


def set_bidi(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = p._p.get_or_add_pPr()
    bidi = ppr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        ppr.append(bidi)
    bidi.set(qn('w:val'), '1')
    for run in p.runs:
        rpr = run._r.get_or_add_rPr()
        rtl = rpr.find(qn('w:rtl'))
        if rtl is None:
            rtl = OxmlElement('w:rtl')
            rpr.append(rtl)
        rtl.set(qn('w:val'), '1')
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.insert(0, fonts)
        for key in ('ascii', 'hAnsi', 'cs'):
            fonts.set(qn('w:' + key), 'Arial')


def field(p, instruction: str):
    run = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = instruction
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t'); placeholder.text = 'حدّث الفهرس في Word لرؤية أرقام الصفحات' if instruction.startswith('TOC') else '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, code, sep, placeholder, end): run._r.append(el)


def docx_build(lines):
    d = Document()
    sec = d.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)
    normal = d.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for style in ('Heading 1', 'Heading 2', 'Title'):
        d.styles[style].font.name = 'Arial'
    d.core_properties.title = 'دليل استخدام نظام كيان ERP — مسودة مصورة جزئياً'
    settings = d.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
    footer = sec.footer.paragraphs[0]
    footer.add_run('دليل كيان — مسودة جزئية  |  صفحة ')
    field(footer, 'PAGE')
    set_bidi(footer)
    for text, style in [('دليل استخدام نظام كيان ERP', 'Title'), ('مسودة عربية مصوّرة جزئيًا — للمراجعة', 'Subtitle'),
                        ('الصور محفوظة من واجهة سابقة. لم تُختبر خطوات البرنامج عمليًا في هذا العمل.', 'Normal'),
                        ('المرجع: شجرة العمل الحالية بالإضافة إلى تغييرات محلية غير مدمجة. HEAD: ' + BASE_SHA, 'Normal')]:
        p=d.add_paragraph(text, style=style); set_bidi(p)
    d.add_page_break()
    p=d.add_paragraph('فهرس الأقسام', style='Title'); set_bidi(p)
    p=d.add_paragraph('قائمة الأقسام أدناه للقراءة فورًا؛ افتح الملف في Word وحدّث الفهرس الآلي لاحتساب أرقام الصفحات.'); set_bidi(p)
    sections, _ = parse_source()
    for number, title in sections:
        p = d.add_paragraph(f'{number}. {title}')
        p.paragraph_format.space_after = Pt(2)
        set_bidi(p)
    p=d.add_paragraph(); field(p, 'TOC \\o "1-1" \\h \\z \\u'); set_bidi(p)
    for raw0 in lines:
        raw=raw0.strip()
        if not raw or raw == '---': continue
        image=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', raw)
        if image:
            p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_path(image.group(2))), width=Cm(16.5))
            continue
        h=re.match(r'^## (\d+)\. (.+)', raw)
        if h:
            p=d.add_paragraph(f'{h.group(1)}. {h.group(2)}', style='Heading 1'); p.paragraph_format.page_break_before=True; set_bidi(p); continue
        if raw.startswith('### '):
            p=d.add_paragraph(clean(raw[4:]), style='Heading 2'); set_bidi(p); continue
        if re.match(r'^\|[\s:|-]+\|$',raw): continue
        if raw.startswith('|'):
            raw=' — '.join(x.strip() for x in raw.strip('|').split('|') if x.strip())
        elif raw.startswith('>'): raw=raw[1:].strip()
        p=d.add_paragraph(clean(raw)); set_bidi(p)
    d.save(DOCX)


def main():
    if not Path(FONT).is_file() or not Path(FONT_BOLD).is_file():
        raise SystemExit('DejaVu Sans fonts are required for Arabic PDF')
    if PDF.exists() or DOCX.exists():
        raise SystemExit('Output exists; refusing to overwrite a deliverable')
    pdfmetrics.registerFont(TTFont('Kayan', FONT))
    pdfmetrics.registerFont(TTFont('KayanBold', FONT_BOLD))
    sections, lines = parse_source()
    OUT.mkdir(parents=True, exist_ok=True)
    measure = GuideDoc(io.BytesIO())
    measure.build(pdf_story(sections, lines, {}))
    if sorted(measure.chapters) != list(range(1,21)):
        raise SystemExit('PDF pagination pass did not find every section')
    pdf = GuideDoc(str(PDF))
    pdf.build(pdf_story(sections, lines, measure.chapters))
    if pdf.chapters != measure.chapters:
        raise SystemExit('PDF index pagination changed during final pass')
    docx_build(lines)
    print('Created:', PDF, DOCX, sep='\n')
    print('PDF chapter pages:', pdf.chapters)


if __name__ == '__main__':
    main()
